from __future__ import annotations

import json
import logging
import os
import re
import unicodedata
import csv
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from dotenv import load_dotenv
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
load_dotenv(ROOT / ".env")


def configure_logging() -> None:
    level = getattr(logging, os.getenv("LOG_LEVEL", "INFO").upper(), logging.INFO)
    work_dir = Path(os.getenv("WORK_DIR", ROOT / "runtime"))
    log_dir = work_dir / "logs"
    log_dir.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        level=level,
        format="%(asctime)s %(levelname)s %(name)s: %(message)s",
        handlers=[
            logging.StreamHandler(),
            logging.FileHandler(log_dir / "oracle-automation.log", encoding="utf-8"),
        ],
        force=True,
    )


def env_bool(name: str, default: bool = False) -> bool:
    return os.getenv(name, str(default)).strip().lower() in {"1", "true", "yes", "sim", "on"}


def normalize_text(value: Any) -> str:
    text = str(value or "").strip()
    return unicodedata.normalize("NFD", text).encode("ascii", "ignore").decode("ascii").upper()


def normalize_dt(value: Any) -> str:
    raw = str(value or "").strip()
    raw = re.sub(r"\.0+$", "", raw)
    digits = re.sub(r"\D", "", raw)
    return digits.lstrip("0") or ("0" if digits else "")


def format_br_datetime(value: datetime | None) -> str:
    return value.strftime("%d/%m/%Y, %H:%M") if value else ""


def format_data_ref(value: date | datetime) -> str:
    return value.strftime("%d/%m/%Y")


def parse_number(value: Any) -> float:
    raw = str(value or "").strip().replace(" ", "")
    if not raw:
        return 0.0
    if "," in raw and "." in raw:
        raw = raw.replace(".", "").replace(",", ".")
    elif "," in raw:
        raw = raw.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:\.\d{3})+", raw):
        raw = raw.replace(".", "")
    raw = re.sub(r"[^\d.-]", "", raw)
    try:
        return float(raw)
    except ValueError:
        return 0.0


def chunks(items: list[Any], size: int) -> Iterable[list[Any]]:
    for index in range(0, len(items), size):
        yield items[index:index + size]


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def read_tabular(path: Path) -> list[list[Any]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return rows
    if suffix == ".xls":
        raise ValueError("Formato .xls antigo nao suportado na VM. Exporte como .xlsx ou .csv.")

    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError(f"Nao foi possivel decodificar {path.name}.")
    sample = "\n".join(text.splitlines()[:10])
    delimiter = "\t" if "\t" in sample else (";" if ";" in sample else ",")
    return [row for row in csv.reader(text.splitlines(), delimiter=delimiter) if any(str(cell).strip() for cell in row)]
